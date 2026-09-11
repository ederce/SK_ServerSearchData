#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Constantes y helpers de estilo Izipay para reportes Excel (openpyxl) -- la
pareja en Excel de la libreria HTML `izipay_report.py` (skill
`reportes-html-izipay`). Misma paleta de marca (rojo #A6342E atenuado) y la
misma semantica de 5 niveles de severidad (good/warning/serious/critical/
error) que usa el HTML, para que un mismo reporte se vea consistente entre
sus dos formatos de salida.

A diferencia de `ReportBuilder` (HTML), esta libreria NO arma la hoja
completa por si sola: cada reporte Excel real tiene una estructura de datos
distinta (columnas, cantidad de hojas, bloques de resumen), asi que expone
**funciones de estilo** para aplicar sobre un `Workbook`/`Worksheet` que el
script llamador ya crea con openpyxl normalmente.

Uso tipico:
    import openpyxl
    from izipay_excel import (title_font, subtitle_font, header_font,
                               header_fill, style_header_row, style_status_cell,
                               set_column_widths, freeze_header)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Titulo del reporte"; ws["A1"].font = title_font()
    ws["A2"] = "Fecha: ... | Comercio: ..."; ws["A2"].font = subtitle_font()
    ws.merge_cells("A1:F1"); ws.merge_cells("A2:F2")

    row = 4
    style_header_row(ws, row, ["Servidor", "Intentos"])
    row += 1
    ws.cell(row, 1, "SVPRDLW364"); ws.cell(row, 2, 1567)

    ws2 = wb.create_sheet("Detalle")
    style_header_row(ws2, 1, ["URL", "Categoria"])
    cat_cell = ws2.cell(2, 2)
    style_status_cell(cat_cell, "critical", label="Proxy Izipay (conocido)")
    set_column_widths(ws2, [55, 22])
    freeze_header(ws2)

    wb.save("reporte.xlsx")
"""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Misma paleta que izipay_report.py (CSS --brand / --status-*), en hex sin
# el "#" (formato que espera openpyxl PatternFill/Font).
BRAND = "A6342E"
BRAND_SOFT = "F7E6E4"
BRAND_INK = "FFFFFF"
INK = "201A1A"
INK_SOFT = "56504F"
INK_MUTE = "948C8B"

STATUS_HEX = {
    "good":     {"fg": "17845A", "soft": "E2F5EC"},
    "warning":  {"fg": "A8710A", "soft": "FAEDD2"},
    "serious":  {"fg": "B3541E", "soft": "FBE3D4"},
    "critical": {"fg": "B3261E", "soft": "FBE7E5"},
    "error":    {"fg": "7A1015", "soft": "F6DCDC"},
    "none":     {"fg": INK_MUTE, "soft": "F4F0ED"},
}


def title_font(size=14):
    return Font(bold=True, size=size, color=BRAND)


def subtitle_font(size=10):
    return Font(size=size, color=INK_SOFT)


def header_font():
    return Font(bold=True, color=BRAND)


def header_fill():
    return PatternFill("solid", fgColor=BRAND_SOFT)


def status_fill(key):
    return PatternFill("solid", fgColor=STATUS_HEX.get(key, STATUS_HEX["none"])["soft"])


def status_font(key, bold=True):
    return Font(bold=bold, color=STATUS_HEX.get(key, STATUS_HEX["none"])["fg"])


def write_title_block(ws, title, subtitle=None, start_row=1, start_col=1, merge_cols=6):
    """Escribe titulo (fila start_row) + subtitulo opcional (fila start_row+1),
    cada uno mergeado start_col..start_col+merge_cols-1. Devuelve la fila
    siguiente libre (start_row+2 si hay subtitulo, start_row+1 si no)."""
    r = start_row
    c1 = ws.cell(r, start_col, title)
    c1.font = title_font()
    if merge_cols > 1:
        ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=start_col + merge_cols - 1)
    r += 1
    if subtitle:
        c2 = ws.cell(r, start_col, subtitle)
        c2.font = subtitle_font()
        if merge_cols > 1:
            ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=start_col + merge_cols - 1)
        r += 1
    return r


def style_header_row(ws, row, headers, start_col=1):
    """Escribe una fila de encabezados de tabla con el estilo de marca
    (fondo brand-soft, texto brand). Devuelve la fila siguiente (row+1)."""
    for i, h in enumerate(headers, start=start_col):
        cell = ws.cell(row, i, h)
        cell.font = header_font()
        cell.fill = header_fill()
    return row + 1


def style_status_cell(cell, key, label=None):
    """Colorea una celda existente segun nivel de severidad (good/warning/
    serious/critical/error/none), mismos colores que los chips del HTML.
    Si se da `label`, tambien fija el valor de la celda."""
    if label is not None:
        cell.value = label
    cell.fill = status_fill(key)
    cell.font = status_font(key)


def set_column_widths(ws, widths, start_col=1):
    for i, w in enumerate(widths, start=start_col):
        ws.column_dimensions[get_column_letter(i)].width = w


def freeze_header(ws, row=2):
    """Congela las filas de encabezado (todo lo que esta arriba de `row`)."""
    ws.freeze_panes = f"A{row}"
